import io
import json
import tempfile
import unittest
from pathlib import Path
from datetime import datetime, timezone
from unittest.mock import patch
import httpx
from openpyxl import load_workbook
from app import purview_analytics as p, usage_directory as d, workbook_reporting as w, claude_analytics as c
from app.compliance_http import ComplianceGate
from app.executive_reporting import executive_xlsx
from app.usage_reporting import usage_csv, usage_html

START=datetime(2026,8,1,tzinfo=timezone.utc)
END=datetime(2026,9,1,tzinfo=timezone.utc)

def event(rid='one',user='a@example.test',when='2026-08-01T12:00:00Z',app='OutlookSidepane'):
    return {'Id':rid,'UserId':user,'CreationTime':when,'Operation':'CopilotInteraction',
            'CopilotEventData':{'AppHost':app,'TargetAgentName':'Draft agent','AccessedResources':[{'Name':'DO NOT RETAIN','Url':'https://private'}]},'Prompt':'DO NOT RETAIN'}

class WorkbookAnalyticsTests(unittest.TestCase):
    def test_audit_metadata_only_and_month_boundaries(self):
        raw=event();row=p.normalize({'auditData':json.dumps(raw)})
        self.assertNotIn('DO NOT RETAIN',json.dumps(row))
        result=p.aggregate([row,row,p.normalize(event('two','b@example.test',app='Office')),p.normalize(event('outside',when='2026-09-01T00:00:00Z'))],START,END)
        self.assertEqual(result['summary']['copilot_interactions'],2)
        self.assertEqual(result['summary']['copilot_active_users'],2)
        self.assertEqual(result['top_users'][0]['active_days'],1)
        self.assertEqual(len(result['copilot_daily']),31)
        self.assertEqual(result['copilot_daily'][1]['Interactions'],0)
        self.assertEqual(result['copilot_detail'][0]['license_type'],'Unknown')
        self.assertEqual(result['copilot_detail'][0]['resources'],1)
        self.assertEqual(result['summary']['agent_interactions'],2)
        result['executive_sections']=w.sections(result)
        names={r['name'] for r in result['executive_sections']}
        self.assertTrue({'Copilot User by App','Copilot Detail','Copilot Daily Trend','Copilot Agents','Copilot App Totals'}<=names)
        book=load_workbook(io.BytesIO(executive_xlsx(result)))
        self.assertEqual(book['Copilot Detail'].max_row,3);book.close()
        self.assertIn('Copilot Detail',usage_csv(result).decode('utf-8-sig'))
        self.assertNotIn('<h2>Copilot Detail</h2>',usage_html(result,'tester','test'))

    def test_conflicting_duplicate_rejected(self):
        with self.assertRaises(ValueError):p.aggregate([p.normalize(event()),p.normalize(event(app='Word'))],START,END)

    def test_usage_and_cost_join_without_counting_cost_as_requests(self):
        actor={'user_id':'one','email':'a@example.test'}
        usage=[{'actor':actor,'product':'chat','model':'model','requests':2,'uncached_input_tokens':10,'cache_read_input_tokens':20,'cache_creation':{'a':5},'output_tokens':4}]
        cost=[{'actor':actor,'product':'chat','model':'model','amount':'125','list_amount':'150','requests':999}]
        details=c.user_detail(usage,cost);users=c.user_rows(details)
        self.assertEqual(users[0]['volume'],2)
        self.assertEqual(users[0]['spend'],1.25)
        self.assertEqual(users[0]['tokens'],39)
        self.assertEqual(details[0]['gross_spend'],1.5)
        self.assertEqual(users[0]['products'],['chat'])

    def test_directory_dated_mapping_includes_zero_usage_departments(self):
        with tempfile.TemporaryDirectory() as tmp,patch.object(p.usage_reporting,'DB_PATH',Path(tmp)/'db'):
            d.save([{'id':'1','userPrincipalName':'a@example.test','mail':'alias@example.test','department':'IT','accountEnabled':True},
                    {'id':'2','userPrincipalName':'b@example.test','department':'HR','accountEnabled':True},
                    {'id':'3','department':'Disabled','accountEnabled':False}], '2026-09-02T00:00:00+00:00')
            data=p.aggregate([p.normalize(event(user='alias@example.test'))],START,END)
            result=d.enrich(data)
            self.assertEqual(result['directory_as_of'],'2026-09-02T00:00:00+00:00')
            self.assertEqual(result['top_users'][0]['department'],'IT')
            rows={r['Department']:r for r in result['department_adoption']}
            self.assertEqual(rows['IT']['Adoption'],1)
            self.assertEqual(rows['HR']['Users'],0)
            self.assertNotIn('Disabled',rows)

class PurviewCollectionTests(unittest.IsolatedAsyncioTestCase):
    async def test_saved_query_resume_paginate_and_publish(self):
        calls=[]
        def respond(request):
            calls.append(request)
            if request.method=='POST':
                self.assertEqual(json.loads(request.content)['operationFilters'],['CopilotInteraction'])
                return httpx.Response(201,json={'id':'query'})
            if request.url.path.endswith('/records'):
                if request.url.params.get('page')=='two':return httpx.Response(200,json={'value':[{'auditData':event('two')}]})
                return httpx.Response(200,json={'value':[{'auditData':event()}],'@odata.nextLink':p.GRAPH+'/security/auditLog/queries/query/records?page=two'})
            return httpx.Response(200,json={'status':'succeeded'})
        with tempfile.TemporaryDirectory() as tmp,patch.object(p.usage_reporting,'DB_PATH',Path(tmp)/'db'),patch.object(p,'compliance_gate',return_value=ComplianceGate(interval=0)):
            async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
                self.assertFalse(await p.collect_step(client,'secret',START,END))
                self.assertIsNone(p.saved('2026-08'))
                self.assertTrue(await p.collect_step(client,'secret',START,END))
                self.assertEqual(p.saved('2026-08')['summary']['copilot_interactions'],2)
                self.assertEqual(sum(r.method=='POST' for r in calls),1)

    async def test_pagination_does_not_send_credentials_to_other_host(self):
        calls=[]
        def respond(request):
            calls.append(request)
            return httpx.Response(200,json={'value':[],'@odata.nextLink':'https://evil.test/records'})
        async with httpx.AsyncClient(transport=httpx.MockTransport(respond)) as client:
            with patch.object(p,'compliance_gate',return_value=ComplianceGate(interval=0)):
                with self.assertRaises(ValueError):await p.graph_pages(client,'secret',p.GRAPH+'/users','/v1.0/users')
        self.assertEqual(len(calls),1)

class PurviewDiagnosticTests(unittest.TestCase):
    def test_permission_failure_includes_stage_and_request_id(self):
        request=httpx.Request('GET','https://graph.microsoft.com/v1.0/security/auditLog/queries/q')
        response=httpx.Response(403,request=request,json={'error':{'code':'Forbidden','message':'Permission denied'}},headers={'request-id':'diagnostic-id'})
        with patch.dict(p.status,{'stage':'check audit query','period':'2026-08'}):
            message=p.http_failure(httpx.HTTPStatusError('failure',request=request,response=response))
        self.assertIn('check audit query',message)
        self.assertIn('2026-08',message)
        self.assertIn('diagnostic-id',message)
        self.assertIn('AuditLogsQuery.Read.All',message)

    def test_authentication_failure_is_not_reported_as_consent_failure(self):
        request=httpx.Request('POST','https://login.microsoftonline.com/tenant/oauth2/v2.0/token')
        response=httpx.Response(401,request=request)
        message=p.http_failure(httpx.HTTPStatusError('failure',request=request,response=response))
        self.assertIn('configured Microsoft credentials',message)
        self.assertNotIn('AuditLogsQuery.Read.All',message)


class SummaryReportTests(unittest.TestCase):
    def test_summary_excludes_raw_records_and_formats_departments(self):
        from app.usage_reporting import summary_tables
        data={'summary':{'copilot_interactions':5000,'copilot_active_users':1},
            'executive_sections':[{'name':'Copilot Detail','rows':[['User'],*[[f'raw-{i}'] for i in range(5000)]]}],
            'department_adoption':[{'Department':'Sales','Headcount':10,'Users':1,'Adoption':.1,'Volume':5000,'Spend (USD)':None}]}
        rendered=usage_html(data,'tester','test')
        self.assertNotIn('raw-4999',rendered)
        self.assertNotIn('Claude',rendered)
        self.assertIn('Department summary',rendered)
        self.assertIn('10.0%',rendered)
        self.assertLess(len(rendered),10000)
        self.assertNotIn('Spend (USD)',str(summary_tables(data)))
