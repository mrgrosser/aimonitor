import io
import unittest
from openpyxl import Workbook, load_workbook
from app.executive_reporting import executive_xlsx
from app.usage_reporting import parse_usage_file, usage_pdf


def sample():
    w=Workbook();s=w.active;s.title="AI Usage Summary"
    for row in [["Monthly report"],["August 2026"],[],["Measure","Copilot","Claude"],["Users with recorded activity",1,1],["Total volume",10,20],["Usage cost in period","Included",3.5],["Agent-assisted volume",2,0],["Distinct surfaces / products",1,1]]:s.append(row)
    sheets={"Department Summary":[["Department","Headcount","Copilot users","Copilot adoption","Claude users"],["IT",2,1,.5,1]],
    "Copilot User by App":[["User","Department","Word","Total","Apps used","Active days"],["person@example.com","IT",10,10,1,2]],
    "Claude User by Product":[["User","Department","Chat","Total requests","Net spend (USD)"],["person@example.com","IT",20,20,3.5]],
    "Claude Detail":[["User","Product","Model","Requests","Net spend (USD)"],["person@example.com","Chat","Test model",20,3.5]],
    "Claude Product & Model":[["Product","Requests","Spend (USD)"],["Chat",20,3.5],["Total",20,3.5],["Model","Requests","Spend (USD)"],["Test model",20,3.5],["Total",20,3.5]],
    "Copilot App Totals":[["App","Interactions","Users"],["Word",10,1]],"Notes & Caveats":[["Different units of measure"]]}
    for name,rows in sheets.items():
        sheet=w.create_sheet(name)
        for row in rows:sheet.append(row)
    w["Department Summary"]["D2"].number_format="0.0%"
    out=io.BytesIO();w.save(out);return out.getvalue()


class ExecutiveReportingTests(unittest.TestCase):
    def test_monthly_totals_and_exact_spend(self):
        data,_=parse_usage_file(sample(),"sample.xlsx")
        self.assertEqual(data["summary"]["claude_usage_spend"],3.5)
        self.assertEqual(data["licensing"],{})
        users=data["top_users"]
        self.assertEqual(len(users),2)
        claude=next(r for r in users if r["provider"]=="Claude Enterprise")
        self.assertEqual(claude["breakdown"][0]["spend"],3.5)
        self.assertNotIn("Department",claude["products"])
        self.assertEqual(data["validation"]["warnings"],[])
        self.assertTrue(usage_pdf(data,"tester","test").startswith(b"%PDF"))

    def test_mismatch_is_rejected(self):
        w=load_workbook(io.BytesIO(sample()));w["AI Usage Summary"]["B6"]=99
        out=io.BytesIO();w.save(out)
        with self.assertRaisesRegex(ValueError,"reconcile"):parse_usage_file(out.getvalue(),"sample.xlsx")

    def test_excel_preserves_sections_formats_and_neutralizes_formulas(self):
        data,_=parse_usage_file(sample(),"sample.xlsx")
        data["executive_sections"][-1]["rows"].append(['=HYPERLINK("https://example.com")'])
        result=load_workbook(io.BytesIO(executive_xlsx(data)))
        self.assertEqual(len(result.sheetnames),8)
        self.assertEqual(result["Department Summary"]["D2"].number_format,"0.0%")
        self.assertEqual(result["Notes & Caveats"]["A2"].data_type,"s")


class ReportChartTests(unittest.TestCase):
    def test_charts_group_small_slices_preserve_totals_and_escape_labels(self):
        from app.report_charts import series,charts_html
        data={"copilot_apps":[{"name":"<script>x</script>"+str(i),"interactions":i+1} for i in range(9)],"claude_products":[],"copilot_daily":[{"Date":"2026-08-02","Interactions":0},{"Date":"2026-08-01","Interactions":5}]}
        apps,daily,_=series(data)
        self.assertEqual(len(apps),6)
        self.assertEqual(sum(v for _,v in apps),45)
        self.assertEqual(daily[0][0],"2026-08-01")
        self.assertNotIn("<script>",charts_html(data))
        self.assertIn("Other apps",charts_html(data))
        self.assertEqual(charts_html({}),"")

    def test_excel_contains_source_linked_native_charts(self):
        data,_=parse_usage_file(sample(),"sample.xlsx")
        book=load_workbook(io.BytesIO(executive_xlsx(data)))
        self.assertEqual(len(book["Copilot App Totals"]._charts),1)
        self.assertEqual(len(book["Claude Product & Model"]._charts),1)
        self.assertIn("'Claude Product & Model'!",book["Claude Product & Model"]._charts[0].ser[0].val.numRef.f)
