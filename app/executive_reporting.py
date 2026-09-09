"""Monthly executive report import and Excel export (separate from evidence scoring)."""
import io
from datetime import date, datetime, timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def import_monthly(workbook, filename, alias):
    def rows(name):
        return [list(r) for r in workbook[name].values] if name in workbook.sheetnames else []
    def table(name, header):
        data=rows(name)
        start=next((i for i,r in enumerate(data) if r and r[0]==header),None)
        if start is None: return [],[]
        columns=data[start]; result=[]
        for row in data[start+1:]:
            if not row or row[0] is None: continue
            if str(row[0]).strip().lower() in {"total","grand total"} or str(row[0]).strip().lower().startswith("total "): break
            result.append(dict(zip(columns,row)))
        return columns,result
    summary_rows=rows("AI Usage Summary")
    import re
    heading=" ".join(str(x or "") for r in summary_rows[:4] for x in r)
    match=re.search(r"(?:January|February|March|April|May|June|July|August|September|October|November|December) 20\d{2}",heading,re.I)
    if not match: raise ValueError("The summary must identify a reporting month and year")
    measures={r[0]:r for r in summary_rows if r and r[0]}
    def measure(name,index):
        row=measures.get(name,[])
        if len(row)<=index or not isinstance(row[index],(int,float)):
            raise ValueError(f"Missing numeric summary measure: {name}")
        return row[index]
    _,apps=table("Copilot App Totals","App")
    _,departments=table("Department Summary","Department")
    _,agents=table("Copilot Agents","Agent")
    _,daily=table("Copilot Daily Trend","Date")
    ch,cusers=table("Copilot User by App","User")
    ah,ausers=table("Claude User by Product","User")
    _,detail=table("Claude Detail","User")
    _,products=table("Claude Product & Model","Product")
    _,models=table("Claude Product & Model","Model")
    top=[]
    def number(value): return value if isinstance(value,(int,float)) else 0
    for headers,users,provider,total,stop in [(ch,cusers,"Microsoft 365 Copilot","Total","Total"),(ah,ausers,"Claude Enterprise","Total requests","Total requests")]:
        cols=[h for h in headers[:headers.index(stop)] if h not in {"User","Department",None}]
        for row in users:
            identity=str(row["User"]); breakdown=[]
            for col in cols:
                volume=number(row.get(col))
                if not volume: continue
                spend=None
                if provider=="Claude Enterprise":
                    matching=[r for r in detail if str(r.get("User","")).casefold()==identity.casefold() and r.get("Product")==col]
                    if matching: spend=round(sum(number(r.get("Net spend (USD)")) for r in matching),2)
                breakdown.append({"name":col,"volume":volume,"spend":spend,"basis":"Imported monthly provider data"})
            top.append({"user":identity,"alias":alias(identity),"department":row.get("Department"),"provider":provider,
                "volume":number(row.get(total)),"surfaces":len(breakdown),"active_days":row.get("Active days"),
                "spend":row.get("Net spend (USD)"),"products":[r["name"] for r in breakdown],"breakdown":breakdown,
                "breakdown_note":"Monthly source data. Copilot interactions and Claude requests are different units."})
    def clean(v): return v.isoformat() if isinstance(v,(date,datetime)) else v
    sections=[]
    for sheet in workbook:
        values=[[clean(v) for v in r] for r in sheet.values]
        sections.append({"name":sheet.title,"rows":values,"formats":[[cell.number_format for cell in row] for row in sheet]})
    data={"period":match.group(0).title(),"source":f"Imported from {filename}","report_schema":2,
        "summary":{"copilot_active_users":measure("Users with recorded activity",1),"claude_active_users":measure("Users with recorded activity",2),
        "copilot_interactions":measure("Total volume",1),"claude_requests":measure("Total volume",2),
        "claude_usage_spend":measure("Usage cost in period",2),"agent_interactions":measure("Agent-assisted volume",1),
        "distinct_surfaces":measure("Distinct surfaces / products",1)},"licensing":{},
        "copilot_apps":[{"name":r["App"],"interactions":r["Interactions"],"users":r["Users"]} for r in apps],
        "claude_products":[{"name":r["Product"],"requests":r["Requests"],"spend":r["Spend (USD)"]} for r in products],
        "claude_models":[{"name":r["Model"],"requests":r["Requests"],"spend":r["Spend (USD)"]} for r in models],
        "departments":departments,"copilot_agents":agents,"copilot_daily":[{k:clean(v) for k,v in r.items()} for r in daily],
        "top_users":sorted(top,key=lambda r:r["volume"],reverse=True),"executive_sections":sections,
        "caveats":[str(r[0]) for r in rows("Notes & Caveats") if r and r[0]]}
    # Reject mismatches rather than publish contradictory executive totals.
    checks=[("Copilot user totals",sum(number(r.get("Total")) for r in cusers),data["summary"]["copilot_interactions"]),
        ("Claude user totals",sum(number(r.get("Total requests")) for r in ausers),data["summary"]["claude_requests"]),
        ("Claude detail spend",round(sum(number(r.get("Net spend (USD)")) for r in detail),2),data["summary"]["claude_usage_spend"])]
    for label,actual,expected in checks:
        if abs(actual-expected)>.01: raise ValueError(f"{label} does not reconcile: {actual} versus {expected}")
    return data


def executive_xlsx(data):
    workbook=Workbook(); workbook.remove(workbook.active)
    sections=data.get("executive_sections") or [{"name":"AI Usage Summary","rows":[["Measure","Value"],*[[k,v] for k,v in data.get("summary",{}).items()]]}]
    for section in sections:
        sheet=workbook.create_sheet(section["name"][:31])
        for row in section["rows"]:
            sheet.append(row)
        for row in sheet:
            populated=[cell for cell in row if cell.value is not None]
            is_heading=populated and (row[0].row==1 or (len(populated)>1 and all(isinstance(c.value,str) for c in populated) and row[0].row<8))
            for cell in row:
                formats=section.get("formats",[])
                if cell.row<=len(formats) and cell.column<=len(formats[cell.row-1]): cell.number_format=formats[cell.row-1][cell.column-1] or "General"
                if isinstance(cell.value,str):
                    try:
                        if len(cell.value)>=19 and cell.value[10]=="T":
                            value=datetime.fromisoformat(cell.value)
                            cell.value=value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
                            cell.number_format='yyyy-mm-dd hh:mm:ss "UTC"'
                    except ValueError: pass
                    if isinstance(cell.value,str): cell.data_type="s"  # Never execute imported formulas.
                cell.font=Font(name="Calibri",size=11,color="FFFFFF" if is_heading else "172231",bold=bool(is_heading))
                cell.alignment=Alignment(vertical="top",wrap_text=True)
                if is_heading: cell.fill=PatternFill("solid",fgColor="243F65")
                elif row[0].row%2==0: cell.fill=PatternFill("solid",fgColor="F2F5F9")
            if len(populated)==1 and sheet.max_column>1:
                sheet.merge_cells(start_row=row[0].row,start_column=1,end_row=row[0].row,end_column=sheet.max_column)
                sheet.row_dimensions[row[0].row].height=max(30,15*(len(str(populated[0].value))//100+1))
        for col in range(1,sheet.max_column+1): sheet.column_dimensions[get_column_letter(col)].width=36 if col==1 else 20
        sheet.freeze_panes="C6" if sheet.max_column>10 else "B5"
        sheet.sheet_view.showGridLines=False
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation="landscape"; sheet.page_setup.paperSize=sheet.PAPERSIZE_A3 if sheet.max_column>10 else sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0
    from app.report_charts import excel_charts
    excel_charts(workbook)
    out=io.BytesIO(); workbook.save(out); return out.getvalue()
